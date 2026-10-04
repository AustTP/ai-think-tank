"""Cluster E line-coverage tests for serve.py.

Covers the library download/ingest/promote/reject/search/passport cluster,
write/read library files, agent file listing/reading, browse + _do_fetch,
allowlist_request, sandbox download/save/curl, curl request sync,
credential/handle management (add_credential, delete_credential,
list_credentials, add_handle, resolve_capability_handle), access_request,
_classify_command, and the room-gate helpers.

Same isolation contract as tests/test_serve.py: every real DB / library /
passport path is redirected into a throwaway temp dir. On top of the standard
boilerplate, a second module patcher also redirects the module-level DERIVED
paths (LIBRARY_ARCHIVE_DIR, LIBRARY_USAGE_PATH, SANDBOXES_DIR,
ESCALATIONS_PATH, BROWSE_TRAIL_PATH) that were computed at import time from the
REAL THINK_TANK_DIR -- without that, list_library/record_library_read/
create_escalation/_sandbox_dir_for would touch the real ~/ai-village-template
tree. Network/process seams (urllib, requests, subprocess, Jev quorum/gate)
are patched per test; library file I/O uses real temp-dir files where that is
simpler and more robust.

Run (isolated coverage file):
  cd /Users/poole86/ai-village-template/world
  COVERAGE_FILE=/tmp/cov_gap_E.coverage python3 -m coverage run --source=serve tests/test_serve_gap_E.py
"""
import asyncio
import builtins
import io
import json
import os
import shutil
import sys
import tempfile
import time
import types
import unittest
import unittest.mock
import urllib.request
import zipfile

import openpyxl

from fastapi.testclient import TestClient
from starlette.requests import Request

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import serve

_MODULE_TMP_DIR = None
_MODULE_PATCHER = None
_EXTRA_PATCHER = None
_RATE_PATCHER = None


def setUpModule():
    global _MODULE_TMP_DIR, _MODULE_PATCHER, _EXTRA_PATCHER, _RATE_PATCHER
    _MODULE_TMP_DIR = tempfile.mkdtemp(prefix='think-tank-serve-gap-')
    _MODULE_PATCHER = unittest.mock.patch.multiple(
        serve,
        DB_PATH=os.path.join(_MODULE_TMP_DIR, 'test.db'),
        THINK_TANK_DIR=_MODULE_TMP_DIR,
        AGENTS_DIR=os.path.join(_MODULE_TMP_DIR, 'agents'),
        LIBRARY_DIR=os.path.join(_MODULE_TMP_DIR, 'library'),
        PASSPORT_PATH=os.path.join(_MODULE_TMP_DIR, 'library', '.passport.json'),
        COLAB_STANDBY_ENABLED=False,
    )
    _MODULE_PATCHER.start()
    # The derived paths below were computed at import time from the REAL
    # THINK_TANK_DIR (see the module docstring) -- redirect them too.
    _EXTRA_PATCHER = unittest.mock.patch.multiple(
        serve,
        LIBRARY_ARCHIVE_DIR=os.path.join(_MODULE_TMP_DIR, 'library', 'archive'),
        LIBRARY_USAGE_PATH=os.path.join(_MODULE_TMP_DIR, 'library_usage.json'),
        SANDBOXES_DIR=os.path.join(_MODULE_TMP_DIR, 'sandboxes'),
        ESCALATIONS_PATH=os.path.join(_MODULE_TMP_DIR, 'escalations.json'),
        BROWSE_TRAIL_PATH=os.path.join(_MODULE_TMP_DIR, 'browse_trail.json'),
    )
    _EXTRA_PATCHER.start()
    # Many endpoint tests share the 'ada' agent id, and check_rate_limit is a
    # real 20-calls/60s window that accrues across the WHOLE module run. Patch
    # it open here (check_rate_limit is not itself a cluster-E target); the
    # dedicated 429 tests override it with return_value=False in their own
    # `with` blocks.
    _RATE_PATCHER = unittest.mock.patch.object(serve, 'check_rate_limit', return_value=True)
    _RATE_PATCHER.start()
    # LIBRARY_DIR must exist before any passport chain append -- CredentialVaultEndpoints
    # / CapabilityHandles run early (alphabetical class order) and _append_passport_decision
    # silently swallows the missing-dir write, leaving blocks empty (IndexError).
    os.makedirs(serve.LIBRARY_DIR, exist_ok=True)
    serve.init_db()


def tearDownModule():
    _RATE_PATCHER.stop()
    _EXTRA_PATCHER.stop()
    _MODULE_PATCHER.stop()
    shutil.rmtree(_MODULE_TMP_DIR, ignore_errors=True)


def _uid(prefix):
    return f'{prefix}-{time.time_ns()}'


def _jev(decision='allow', confidence=0.95, trace='trace-gap'):
    return unittest.mock.AsyncMock(return_value=(decision, confidence, 0.001, trace))


class _FakeResp:
    """Minimal urllib response stand-in: context manager with geturl() /
    read(n) / headers / status, enough for _download_file_sync and
    _curl_request_sync."""

    def __init__(self, url, headers, status, data):
        self._url = url
        self.headers = headers
        self.status = status
        self._data = data

    def geturl(self):
        return self._url

    def read(self, n=-1):
        if n is None or n < 0:
            return self._data
        return self._data[:n]

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


def _write_lib(rel_path, content, mtime=None):
    full = os.path.join(serve.LIBRARY_DIR, rel_path)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, 'w') as f:
        f.write(content)
    if mtime is not None:
        os.utime(full, (mtime, mtime))
    return full


def _client():
    return TestClient(serve.app)


class DownloadFileSync(unittest.TestCase):
    def test_success_reads_and_returns_raw(self):
        fake = _FakeResp('https://example.com/data.csv', {'Content-Type': 'text/csv'}, 200, b'a,b\n1,2\n')
        with unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve.urllib.request, 'urlopen', return_value=fake):
            final_url, content_type, raw, truncated = serve._download_file_sync('https://example.com/data.csv', 1000)
        self.assertEqual(final_url, 'https://example.com/data.csv')
        self.assertEqual(content_type, 'text/csv')
        self.assertEqual(raw, b'a,b\n1,2\n')
        self.assertFalse(truncated)

    def test_truncates_at_max_bytes(self):
        fake = _FakeResp('https://example.com/big.bin', {'Content-Type': 'application/octet-stream'}, 200, b'x' * 50)
        with unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve.urllib.request, 'urlopen', return_value=fake):
            _final, _ct, raw, truncated = serve._download_file_sync('https://example.com/big.bin', 10)
        self.assertEqual(len(raw), 10)
        self.assertTrue(truncated)

    def test_redirect_to_disallowed_host_raises(self):
        fake = _FakeResp('http://10.0.0.5/steal', {'Content-Type': 'text/html'}, 200, b'x')
        with unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=False), \
             unittest.mock.patch.object(serve.urllib.request, 'urlopen', return_value=fake):
            with self.assertRaises(ValueError):
                serve._download_file_sync('https://example.com/ok', 100)


class LibraryDownloadEndpoint(unittest.TestCase):
    def _post(self, body, **patches):
        return serve.verify_agent_key, body

    def test_disabled_browsing_returns_403(self):
        with unittest.mock.patch.object(serve, 'BROWSING_ENABLED', False), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/download', json={'url': 'http://x', 'filename': 'x'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('disabled', r.json()['error'])

    def test_missing_api_key_returns_500(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/download', json={'url': 'http://x', 'filename': 'x'})
        self.assertEqual(r.status_code, 500)

    def test_missing_url_or_filename_returns_400(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().post('/api/library/download', json={'url': '', 'filename': ''})
        self.assertEqual(r.status_code, 400)
        self.assertIn('url and filename are required', r.json()['error'])

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/download', json={'url': 'http://x', 'filename': 'x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 429)
        log.assert_called_once()
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_unsafe_filename_returns_400(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().post('/api/library/download', json={'url': 'http://x', 'filename': '///', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('safe characters', r.json()['error'])

    def test_non_http_scheme_is_blocked(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/download', json={'url': 'ftp://example.com/a.csv', 'filename': 'a.csv', 'agentId': 'ada', 'purpose': 'data'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'invalid or non-http(s) scheme')

    def test_private_host_is_blocked(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/download', json={'url': 'http://internal.example/a.csv', 'filename': 'a.csv', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'private/internal host')

    def test_jev_denial_blocks(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False):
            r = _client().post('/api/library/download', json={'url': 'http://example.com/a.csv', 'filename': 'a.csv', 'agentId': 'ada', 'purpose': 'ref'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])

    def test_fetch_failure_reports_ok_false(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync', side_effect=ValueError('boom')), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/download', json={'url': 'http://example.com/a.csv', 'filename': 'a.csv', 'agentId': 'ada', 'purpose': 'ref'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['allowed'])
        self.assertFalse(r.json()['ok'])
        self.assertEqual(log.call_args[0][2]['reason'], 'boom')

    def test_personal_download_success_writes_file(self):
        agent = _uid('dl-agent')
        body = {'url': 'https://example.com/data.csv', 'filename': '../data.csv', 'agentId': agent, 'purpose': 'a dataset'}
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev(trace='trace-dl')), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync',
                                        return_value=('https://example.com/data.csv', 'text/csv', b'a,b\n1,2\n', False)), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/download', json=body)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()['allowed'])
        self.assertTrue(r.json()['ok'])
        rel = r.json()['path']
        self.assertTrue(rel.startswith(f'pending_review/downloads/{agent}/'))
        full = os.path.join(serve.LIBRARY_DIR, rel)
        self.assertTrue(os.path.isfile(full), 'the download must land on disk')
        with open(full, 'rb') as f:
            self.assertEqual(f.read(), b'a,b\n1,2\n')
        self.assertEqual(log.call_args[0][2]['decision'], 'allowed')

    def test_shared_scope_routes_to_shared_pending(self):
        agent = _uid('dl-agent')
        body = {'url': 'https://example.com/data.csv', 'filename': 'data.csv', 'agentId': agent, 'scope': 'shared', 'purpose': 'team data'}
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync',
                                        return_value=('https://example.com/data.csv', 'text/csv', b'x', False)):
            r = _client().post('/api/library/download', json=body)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['path'].startswith('pending_review/shared/'))

    def test_bogus_scope_defaults_to_personal(self):
        agent = _uid('dl-agent')
        body = {'url': 'https://example.com/data.csv', 'filename': 'data.csv', 'agentId': agent, 'scope': 'bogus', 'purpose': 'x'}
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync',
                                        return_value=('https://example.com/data.csv', 'text/csv', b'x', False)):
            r = _client().post('/api/library/download', json=body)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['path'].startswith('pending_review/downloads/'))

    def test_traversal_agent_id_is_rejected(self):
        body = {'url': 'https://example.com/data.csv', 'filename': 'data.csv', 'agentId': '../../..', 'purpose': 'x'}
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync',
                                        return_value=('https://example.com/data.csv', 'text/csv', b'x', False)):
            r = _client().post('/api/library/download', json=body)
        self.assertEqual(r.status_code, 400)
        self.assertIn('invalid destination', r.json()['error'])


class PassportLedger(unittest.TestCase):
    def setUp(self):
        if os.path.exists(serve.PASSPORT_PATH):
            os.remove(serve.PASSPORT_PATH)

    def test_load_missing_file_is_fresh_chain(self):
        self.assertEqual(serve._load_passport(), {'version': 1, 'count': 0, 'head': None, 'blocks': []})

    def test_load_corrupt_file_is_fresh_chain(self):
        os.makedirs(os.path.dirname(serve.PASSPORT_PATH), exist_ok=True)
        with open(serve.PASSPORT_PATH, 'w') as f:
            f.write('not json{{{')
        self.assertEqual(serve._load_passport(), {'version': 1, 'count': 0, 'head': None, 'blocks': []})

    def test_load_backfills_missing_keys(self):
        os.makedirs(os.path.dirname(serve.PASSPORT_PATH), exist_ok=True)
        with open(serve.PASSPORT_PATH, 'w') as f:
            json.dump({'blocks': [{'index': 1}]}, f)
        data = serve._load_passport()
        self.assertEqual(data['version'], 1)
        self.assertEqual(data['count'], 0)
        self.assertIsNone(data['head'])
        self.assertEqual(len(data['blocks']), 1)

    def test_append_skips_missing_file(self):
        self.assertIsNone(serve._append_passport('does/not/exist.md', 'ada', 'player'))

    def test_append_chains_blocks_and_writes_ledger(self):
        path = _write_lib('pending_review/downloads/ada/notes.md', 'real content')
        block = serve._append_passport('pending_review/downloads/ada/notes.md', 'ada', 'player')
        self.assertIsNotNone(block)
        self.assertEqual(block['index'], 1)
        self.assertEqual(block['path'], 'pending_review/downloads/ada/notes.md')
        self.assertEqual(block['sha256'], serve._sha256_file(path))
        self.assertEqual(block['owner'], 'ada')
        self.assertEqual(block['promotedBy'], 'player')
        self.assertIsNone(block['prev'])
        data = serve._load_passport()
        self.assertEqual(data['count'], 1)
        self.assertEqual(data['head'], block['sha256'])
        self.assertEqual(len(data['blocks']), 1)

        second = serve._append_passport('pending_review/downloads/ada/notes.md', 'ada', 'player')
        self.assertEqual(second['index'], 2)
        self.assertEqual(second['prev'], block['sha256'], 'each block links back to the previous head')

    def test_append_failure_returns_none(self):
        with unittest.mock.patch.object(serve, '_safe_library_path', side_effect=RuntimeError('boom')):
            self.assertIsNone(serve._append_passport('shared/x.md', 'ada', 'player'))

    def test_library_passport_endpoint_returns_ledger(self):
        _write_lib('shared/known.md', 'known content')
        serve._append_passport('shared/known.md', 'ada', 'player')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/passport')
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body['count'], 1)
        self.assertEqual(len(body['blocks']), 1)


class IngestPdf(unittest.TestCase):
    def test_missing_pypdf_returns_clear_error(self):
        with unittest.mock.patch.object(serve, 'pypdf', None):
            text, err = serve._ingest_pdf_to_text('/x/whatever.pdf')
        self.assertIsNone(text)
        self.assertIn('pypdf is not installed', err)

    def test_success_extracts_pages(self):
        class FakePage:
            def extract_text(self):
                return 'page body'
        class FakeReader:
            pages = [FakePage(), FakePage()]
        with unittest.mock.patch.object(serve, 'pypdf', types.SimpleNamespace(PdfReader=lambda p: FakeReader())):
            text, err = serve._ingest_pdf_to_text('/x/doc.pdf')
        self.assertIsNone(err)
        self.assertIn('--- page 1 ---', text)
        self.assertIn('page body', text)
        self.assertIn('--- page 2 ---', text)

    def test_more_than_200_pages_is_capped(self):
        class FakePage:
            def extract_text(self):
                return 'p'
        class FakeReader:
            pages = [FakePage() for _ in range(201)]
        with unittest.mock.patch.object(serve, 'pypdf', types.SimpleNamespace(PdfReader=lambda p: FakeReader())):
            text, err = serve._ingest_pdf_to_text('/x/big.pdf')
        self.assertIsNone(err)
        self.assertIn('1 more pages not extracted', text)

    def test_extraction_exception_is_reported(self):
        with unittest.mock.patch.object(serve, 'pypdf', types.SimpleNamespace(PdfReader=lambda p: (_ for _ in ()).throw(ValueError('corrupt pdf')))):
            text, err = serve._ingest_pdf_to_text('/x/bad.pdf')
        self.assertIsNone(text)
        self.assertIn('corrupt pdf', err)


class IngestXlsx(unittest.TestCase):
    def test_missing_openpyxl_returns_clear_error(self):
        with unittest.mock.patch.object(serve, 'openpyxl', None):
            text, err = serve._ingest_xlsx_to_text('/x/whatever.xlsx')
        self.assertIsNone(text)
        self.assertIn('openpyxl is not installed', err)

    def test_success_extracts_sheets_and_rows(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'book.xlsx')
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Data'
            ws.append(['a', 'b'])
            ws.append([1, 2])
            wb.save(path)
            text, err = serve._ingest_xlsx_to_text(path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertIsNone(err)
        self.assertIn('--- sheet: Data ---', text)
        self.assertIn('a, b', text)
        self.assertIn('1, 2', text)

    def test_over_2000_rows_is_truncated(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'wide.xlsx')
            wb = openpyxl.Workbook()
            ws = wb.active
            for i in range(2001):
                ws.append([i])
            wb.save(path)
            text, err = serve._ingest_xlsx_to_text(path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertIsNone(err)
        self.assertIn('truncated at 2000 rows', text)

    def test_load_exception_is_reported(self):
        with unittest.mock.patch.object(serve.openpyxl, 'load_workbook', side_effect=Exception('bad workbook')):
            text, err = serve._ingest_xlsx_to_text('/x/book.xlsx')
        self.assertIsNone(text)
        self.assertIn('bad workbook', err)


class IngestWriteText(unittest.TestCase):
    def test_invalid_destination_reports_failure(self):
        results = []
        serve._write_ingested_text('../../escape.md', 'content', results, 'name.md')
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertEqual(results[0]['note'], 'invalid destination path')

    def test_valid_destination_writes_file(self):
        results = []
        serve._write_ingested_text('ingested/notes/name.md', 'hello world', results, 'name.md')
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]['ok'])
        self.assertEqual(results[0]['path'], 'ingested/notes/name.md')
        with open(os.path.join(serve.LIBRARY_DIR, 'ingested/notes/name.md')) as f:
            self.assertEqual(f.read(), 'hello world')


class IngestOneFile(unittest.TestCase):
    def test_getsize_failure_is_reported(self):
        results = []
        serve._ingest_one_file('/nonexistent/nope.pdf', 'ingested', results)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])

    def test_oversized_file_is_skipped(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'big.txt')
            with open(path, 'w') as f:
                f.write('x' * 20)
            results = []
            with unittest.mock.patch.object(serve, 'INGEST_MAX_FILE_BYTES', 10):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertIn('exceeds the', results[0]['note'])

    def test_pdf_success_writes_markdown(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'doc.pdf')
            with open(path, 'wb') as f:
                f.write(b'%PDF-1.4 fake')
            results = []
            with unittest.mock.patch.object(serve, '_ingest_pdf_to_text', return_value=('extracted pdf text', None)):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]['ok'])
        self.assertEqual(results[0]['path'], 'ingested/doc.pdf.md')
        with open(os.path.join(serve.LIBRARY_DIR, 'ingested/doc.pdf.md')) as f:
            self.assertIn('extracted pdf text', f.read())

    def test_pdf_extraction_error_is_reported(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'bad.pdf')
            with open(path, 'wb') as f:
                f.write(b'garbage')
            results = []
            with unittest.mock.patch.object(serve, '_ingest_pdf_to_text', return_value=(None, 'PDF extraction failed: nope')):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertIn('PDF extraction failed: nope', results[0]['note'])

    def test_xlsx_success_writes_markdown(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'data.xlsx')
            with open(path, 'wb') as f:
                f.write(b'fake xlsx')
            results = []
            with unittest.mock.patch.object(serve, '_ingest_xlsx_to_text', return_value=('sheet rows', None)):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]['ok'])
        self.assertEqual(results[0]['path'], 'ingested/data.xlsx.md')

    def test_xlsx_extraction_error_is_reported(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'bad.xlsx')
            with open(path, 'wb') as f:
                f.write(b'garbage')
            results = []
            with unittest.mock.patch.object(serve, '_ingest_xlsx_to_text', return_value=(None, 'Excel extraction failed: nope')):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertIn('Excel extraction failed: nope', results[0]['note'])

    def test_text_file_success(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'note.txt')
            with open(path, 'w') as f:
                f.write('plain text content')
            results = []
            serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]['ok'])
        with open(os.path.join(serve.LIBRARY_DIR, 'ingested/note.txt')) as f:
            self.assertEqual(f.read(), 'plain text content')

    def test_text_read_error_is_reported(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'note.txt')
            with open(path, 'w') as f:
                f.write('x')
            results = []
            real_open = builtins.open

            def fake_open(p, *a, **k):
                if p == path:
                    raise PermissionError('denied')
                return real_open(p, *a, **k)
            with unittest.mock.patch.object(builtins, 'open', side_effect=fake_open):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertIn('denied', results[0]['note'])

    def test_image_success_is_copied(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'photo.png')
            with open(path, 'wb') as f:
                f.write(b'\x89PNG fake image bytes')
            results = []
            serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]['ok'])
        self.assertIn('vision review', results[0]['note'])
        with open(os.path.join(serve.LIBRARY_DIR, 'ingested/photo.png'), 'rb') as f:
            self.assertEqual(f.read(), b'\x89PNG fake image bytes')

    def test_image_read_error_is_reported(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'photo.png')
            with open(path, 'wb') as f:
                f.write(b'png')
            results = []
            real_open = builtins.open

            def fake_open(p, *a, **k):
                if p == path:
                    raise PermissionError('no read')
                return real_open(p, *a, **k)
            with unittest.mock.patch.object(builtins, 'open', side_effect=fake_open):
                serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertIn('no read', results[0]['note'])

    def test_image_invalid_destination_is_reported(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'photo.png')
            with open(path, 'wb') as f:
                f.write(b'png')
            results = []
            serve._ingest_one_file(path, '../../escape', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertEqual(results[0]['note'], 'invalid destination path')

    def test_unsupported_extension_is_skipped(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'archive.rar')
            with open(path, 'wb') as f:
                f.write(b'rar')
            results = []
            serve._ingest_one_file(path, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]['ok'])
        self.assertIn('unsupported file type', results[0]['note'])


class IngestWalk(unittest.TestCase):
    def test_skips_ds_store_and_ingests_real_files(self):
        tmp = tempfile.mkdtemp()
        try:
            with open(os.path.join(tmp, '.DS_Store'), 'w') as f:
                f.write('junk')
            os.makedirs(os.path.join(tmp, 'sub'))
            with open(os.path.join(tmp, 'sub', 'a.txt'), 'w') as f:
                f.write('a')
            results = []
            serve._ingest_walk(tmp, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]['ok'])
        self.assertEqual(results[0]['path'], 'ingested/sub/a.txt')

    def test_cap_reached_notes_skipped_files(self):
        tmp = tempfile.mkdtemp()
        try:
            for n in ('a.txt', 'b.txt'):
                with open(os.path.join(tmp, n), 'w') as f:
                    f.write(n)
            results = []
            with unittest.mock.patch.object(serve, 'INGEST_MAX_FILES', 1):
                serve._ingest_walk(tmp, 'ingested', results)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(len(results), 2)
        self.assertTrue(any(r['ok'] for r in results))
        self.assertTrue(any(not r['ok'] and 'cap of 1 files reached' in r['note'] for r in results))


class IngestLibraryEndpoint(unittest.TestCase):
    def test_agent_cannot_ingest(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/ingest', json={'sourcePath': '/tmp', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('player-only', r.json()['error'])

    def test_missing_source_path_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/ingest', json={})
        self.assertEqual(r.status_code, 400)

    def test_nonexistent_source_returns_404(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/ingest', json={'sourcePath': '/nonexistent/nowhere'})
        self.assertEqual(r.status_code, 404)

    def test_single_file_ingest(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'note.md')
            with open(path, 'w') as f:
                f.write('# a real note')
            with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
                 unittest.mock.patch.object(serve, 'log_action') as log:
                r = _client().post('/api/library/ingest', json={'sourcePath': path, 'destPath': 'ingested/notes'})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertTrue(body['ok'])
        self.assertEqual(body['okCount'], 1)
        self.assertEqual(body['totalCount'], 1)
        self.assertTrue(os.path.isfile(os.path.join(serve.LIBRARY_DIR, 'ingested/notes/note.md')))
        self.assertEqual(log.call_args[0][1], 'library_ingest')

    def test_directory_walk_ingest(self):
        tmp = tempfile.mkdtemp()
        try:
            os.makedirs(os.path.join(tmp, 'src', 'sub'))
            for rel in ('src/a.txt', 'src/sub/b.txt'):
                with open(os.path.join(tmp, rel), 'w') as f:
                    f.write(rel)
            with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
                r = _client().post('/api/library/ingest', json={'sourcePath': os.path.join(tmp, 'src')})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['okCount'], 2)

    def test_zip_ingest_extracts_and_walks(self):
        tmp = tempfile.mkdtemp()
        try:
            zpath = os.path.join(tmp, 'bundle.zip')
            with zipfile.ZipFile(zpath, 'w') as zf:
                zf.writestr('one.txt', 'one')
                zf.writestr('sub/two.txt', 'two')
            with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
                r = _client().post('/api/library/ingest', json={'sourcePath': zpath})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['okCount'], 2)
        self.assertTrue(os.path.isfile(os.path.join(serve.LIBRARY_DIR, 'ingested/bundle.zip/one.txt')))

    def test_zip_bomb_guard_refuses(self):
        tmp = tempfile.mkdtemp()
        try:
            zpath = os.path.join(tmp, 'bomb.zip')
            with open(zpath, 'wb') as f:
                f.write(b'not really a zip')
            class FakeZF:
                def __enter__(self):
                    return self
                def __exit__(self, *a):
                    return False
                def infolist(self):
                    return [types.SimpleNamespace(file_size=serve.INGEST_MAX_FILE_BYTES * 11)]
                def extractall(self, dest):
                    raise AssertionError('must never extract')
            with unittest.mock.patch.object(serve.zipfile, 'ZipFile', return_value=FakeZF()), \
                 unittest.mock.patch.object(serve, 'verify_session', return_value=True):
                r = _client().post('/api/library/ingest', json={'sourcePath': zpath})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(r.status_code, 400)
        self.assertIn('exceed the ingest cap', r.json()['error'])

    def test_bad_zip_returns_400(self):
        tmp = tempfile.mkdtemp()
        try:
            zpath = os.path.join(tmp, 'notzip.zip')
            with open(zpath, 'wb') as f:
                f.write(b'garbage bytes, not a zip at all')
            with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
                r = _client().post('/api/library/ingest', json={'sourcePath': zpath})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        self.assertEqual(r.status_code, 400)
        self.assertIn('not a valid zip', r.json()['error'])


class WriteLibraryFile(unittest.TestCase):
    def setUp(self):
        if os.path.exists(serve.PASSPORT_PATH):
            os.remove(serve.PASSPORT_PATH)

    def _post(self, body):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            return _client().post('/api/library/file', json=body)

    def test_missing_path_returns_400(self):
        r = self._post({'agentId': 'ada', 'content': 'x'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('path is required', r.json()['error'])

    def test_external_source_forced_into_pending_review(self):
        r = self._post({'agentId': 'ada', 'path': 'shared/snippet.md', 'content': 'browsed text', 'source': 'external'})
        self.assertEqual(r.status_code, 200, r.text)
        full = os.path.join(serve.LIBRARY_DIR, 'pending_review/shared/snippet.md')
        self.assertTrue(os.path.isfile(full))

    def test_invalid_path_returns_400(self):
        r = self._post({'agentId': 'ada', 'path': '../../.env', 'content': 'x'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('invalid path', r.json()['error'])

    def test_wiki_path_is_director_gated(self):
        r = self._post({'agentId': 'ada', 'path': 'wiki/live.md', 'content': 'x'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('director-gated', r.json()['error'])

    def test_working_guide_requires_director(self):
        state = {'agents': {}, 'agentRoster': [{'id': 'sam'}]}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'get_state_from_db', return_value=state):
            r = _client().post('/api/library/file', json={'agentId': 'sam', 'path': 'working-guide.md', 'content': 'x'})
        self.assertEqual(r.status_code, 403)

    def test_working_guide_allowed_for_admin(self):
        state = {'agents': {}, 'agentRoster': [{'id': 'faye', 'isAdmin': True}]}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'get_state_from_db', return_value=state):
            r = _client().post('/api/library/file', json={'agentId': 'faye', 'path': 'working-guide.md', 'content': 'guide'})
        self.assertEqual(r.status_code, 200, r.text)
        with open(os.path.join(serve.LIBRARY_DIR, 'working-guide.md')) as f:
            self.assertEqual(f.read(), 'guide')

    def test_another_agents_personal_dir_is_denied(self):
        state = {'agents': {}, 'agentRoster': []}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'get_state_from_db', return_value=state), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/file', json={'agentId': 'ada', 'path': 'downloads/eli/private.md', 'content': 'x'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('Only eli', r.json()['error'])
        self.assertEqual(log.call_args[0][1], 'library_write_denied')

    def test_own_personal_dir_is_allowed(self):
        r = self._post({'agentId': 'ada', 'path': 'downloads/ada/mine.md', 'content': 'mine'})
        self.assertEqual(r.status_code, 200, r.text)
        with open(os.path.join(serve.LIBRARY_DIR, 'downloads/ada/mine.md')) as f:
            self.assertEqual(f.read(), 'mine')

    def test_success_writes_redacted_and_chains_passport(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/file', json={'agentId': 'ada', 'path': 'shared/idea.md', 'content': 'sk-abcdefghijklmnopqrstuvwxyz123456 idea'})
        self.assertEqual(r.status_code, 200, r.text)
        with open(os.path.join(serve.LIBRARY_DIR, 'shared/idea.md')) as f:
            content = f.read()
        self.assertNotIn('sk-abcdefghijklmnopqrstuvwxyz123456', content)
        self.assertIn('[REDACTED]', content)
        self.assertEqual(log.call_args[0][1], 'library_write')
        data = serve._load_passport()
        self.assertEqual(data['count'], 1)
        self.assertEqual(data['blocks'][0]['kind'], 'library_write')


class LibraryPromote(unittest.TestCase):
    def setUp(self):
        if os.path.exists(serve.PASSPORT_PATH):
            os.remove(serve.PASSPORT_PATH)

    def _post(self, body, **extra):
        patchers = [unittest.mock.patch.object(serve, 'verify_session', return_value=True),
                    unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True)]
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            return _client().post('/api/library/promote', json=body)

    def test_path_outside_pending_review_is_400(self):
        r = self._post({'agentId': 'ada', 'path': 'shared/x.md'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('pending_review/', r.json()['error'])

    def test_missing_source_file_is_404(self):
        r = self._post({'agentId': 'ada', 'path': 'pending_review/nope.md'})
        self.assertEqual(r.status_code, 404)

    def test_empty_destination_is_404(self):
        r = self._post({'agentId': 'ada', 'path': 'pending_review/'})
        self.assertEqual(r.status_code, 404)

    def test_invalid_destination_is_400(self):
        _write_lib('pending_review/x.md', 'x')
        real_safe = serve._safe_library_path

        def fake(p):
            if p == 'pending_review/x.md':
                return real_safe(p)
            return None
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_safe_library_path', side_effect=fake):
            r = _client().post('/api/library/promote', json={'agentId': 'ada', 'path': 'pending_review/x.md'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('invalid destination', r.json()['error'])

    def test_into_another_agents_personal_dir_is_denied(self):
        _write_lib('pending_review/downloads/eli/x.md', 'x')
        state = {'agents': {}, 'agentRoster': []}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'get_state_from_db', return_value=state), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/promote', json={'agentId': 'ada', 'path': 'pending_review/downloads/eli/x.md'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(log.call_args[0][1], 'library_promote_denied')

    def test_into_wiki_is_denied(self):
        _write_lib('pending_review/wiki/x.md', 'x')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/promote', json={'agentId': 'ada', 'path': 'pending_review/wiki/x.md'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('director-gated', r.json()['error'])

    def test_success_moves_file_and_chains_passport(self):
        _write_lib('pending_review/shared/x.md', 'trusted content')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/promote', json={'agentId': 'ada', 'path': 'pending_review/shared/x.md'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.text, 'promoted')
        self.assertFalse(os.path.exists(os.path.join(serve.LIBRARY_DIR, 'pending_review/shared/x.md')))
        with open(os.path.join(serve.LIBRARY_DIR, 'shared/x.md')) as f:
            self.assertEqual(f.read(), 'trusted content')
        self.assertEqual(log.call_args[0][1], 'library_promote')
        data = serve._load_passport()
        kinds = [b.get('kind') for b in data['blocks']]
        self.assertIn(None, kinds, 'a promoted-file block is chained')
        self.assertIn('library_promote', kinds)


class LibraryReject(unittest.TestCase):
    def test_path_outside_pending_review_is_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/reject', json={'agentId': 'ada', 'path': 'shared/x.md'})
        self.assertEqual(r.status_code, 400)

    def test_missing_source_file_is_404(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/library/reject', json={'agentId': 'ada', 'path': 'pending_review/nope.md'})
        self.assertEqual(r.status_code, 404)

    def test_invalid_destination_is_400(self):
        _write_lib('pending_review/x.md', 'x')
        real_safe = serve._safe_library_path

        def fake(p):
            if p == 'pending_review/x.md':
                return real_safe(p)
            return None
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_safe_library_path', side_effect=fake):
            r = _client().post('/api/library/reject', json={'agentId': 'ada', 'path': 'pending_review/x.md'})
        self.assertEqual(r.status_code, 400)

    def test_success_moves_to_rejected(self):
        _write_lib('pending_review/x.md', 'untrusted')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/library/reject', json={'agentId': 'ada', 'path': 'pending_review/x.md'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.text, 'rejected')
        self.assertTrue(os.path.isfile(os.path.join(serve.LIBRARY_DIR, 'rejected/x.md')))
        self.assertFalse(os.path.exists(os.path.join(serve.LIBRARY_DIR, 'pending_review/x.md')))
        self.assertEqual(log.call_args[0][1], 'library_reject')


class ListAndSearchLibrary(unittest.TestCase):
    def test_list_skips_dotfiles_and_sorts_by_modified(self):
        old = _write_lib('shared/old.md', 'old', mtime=time.time() - 1000)
        new = _write_lib('shared/new.md', 'new', mtime=time.time())
        _write_lib('.passport.json', '{}', mtime=time.time())
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().get('/api/library')
        self.assertEqual(r.status_code, 200)
        files = r.json()['files']
        paths = [f['path'] for f in files]
        self.assertIn('shared/old.md', paths)
        self.assertIn('shared/new.md', paths)
        self.assertNotIn('.passport.json', paths)
        self.assertEqual(paths[0], 'shared/new.md', 'newest-modified file sorts first')
        entry = [f for f in files if f['path'] == 'shared/old.md'][0]
        self.assertEqual(entry['size'], os.path.getsize(old))

    def test_search_empty_query_is_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().get('/api/library/search', params={'q': '   '})
        self.assertEqual(r.status_code, 400)
        self.assertIn('q is required', r.json()['error'])

    def test_search_returns_matches(self):
        _write_lib('shared/note.md', 'the think tank studies weather patterns')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().get('/api/library/search', params={'q': 'weather'})
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertEqual(body['query'], 'weather')
        self.assertTrue(any(m['path'] == 'shared/note.md' for m in body['matches']))

    def test_search_matches_empty_query_is_empty(self):
        self.assertEqual(serve._library_search_matches('   '), [])

    def test_search_matches_skips_images(self):
        _write_lib('shared/photo.png', 'contains weather but is an image')
        _write_lib('shared/txt.md', 'the weather report')
        matches = serve._library_search_matches('weather')
        paths = [m['path'] for m in matches]
        self.assertIn('shared/txt.md', paths)
        self.assertNotIn('shared/photo.png', paths)

    def test_search_matches_skips_unreadable_files(self):
        _write_lib('shared/ok.md', 'weather here')
        broken = os.path.join(serve.LIBRARY_DIR, 'shared/broken.md')
        os.symlink('/nonexistent/target-file', broken)
        matches = serve._library_search_matches('weather')
        paths = [m['path'] for m in matches]
        self.assertIn('shared/ok.md', paths)
        self.assertNotIn('shared/broken.md', paths)

    def test_search_matches_skips_no_hit_files(self):
        _write_lib('shared/other.md', 'nothing relevant in here')
        _write_lib('shared/weather.md', 'sunny and warm today')
        matches = serve._library_search_matches('sunny')
        self.assertEqual([m['path'] for m in matches], ['shared/weather.md'])

    def test_search_matches_builds_snippet_and_size(self):
        content = 'the word weather appears ' + ('filler ' * 200)
        _write_lib('shared/big.md', content)
        matches = serve._library_search_matches('weather')
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]['size'], len(content))
        self.assertIn('weather', matches[0]['snippet'])

    def test_search_ranks_trail_reinforced_first(self):
        _write_lib('shared/validated.md', 'weather research', mtime=time.time() - 1000)
        _write_lib('shared/fresh.md', 'weather research', mtime=time.time())
        serve.record_library_read('shared/validated.md')
        serve.record_library_read('shared/validated.md')
        matches = serve._library_search_matches('weather')
        self.assertEqual(matches[0]['path'], 'shared/validated.md')


class ReadLibraryFile(unittest.TestCase):
    def test_missing_file_returns_404(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().get('/api/library/file', params={'path': 'shared/nope.md'})
        self.assertEqual(r.status_code, 404)
        self.assertIn('not found', r.json()['error'])

    def test_success_reads_and_records_usage(self):
        _write_lib('shared/readme.md', 'hello library')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'record_library_read') as record:
            r = _client().get('/api/library/file', params={'path': 'shared/readme.md'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['content'], 'hello library')
        record.assert_called_once_with('shared/readme.md')


class OwnsLibraryPath(unittest.TestCase):
    def test_short_paths_are_commons(self):
        self.assertIsNone(serve._owns_library_path('shared'))
        self.assertIsNone(serve._owns_library_path('x'))

    def test_downloads_map_to_owner(self):
        self.assertEqual(serve._owns_library_path('downloads/eli/file.md'), 'eli')

    def test_pending_review_downloads_map_to_owner(self):
        self.assertEqual(serve._owns_library_path('pending_review/downloads/eli/file.md'), 'eli')

    def test_other_pending_paths_are_commons(self):
        self.assertIsNone(serve._owns_library_path('pending_review/shared/file.md'))


class ResolveRequester(unittest.TestCase):
    def _req(self, query='', headers=None):
        scope = {
            'type': 'http',
            'method': 'GET',
            'path': '/api/agent-files',
            'query_string': query.encode(),
            'headers': headers or [],
        }
        return Request(scope)

    def test_no_claimed_identity_returns_none(self):
        self.assertIsNone(serve._resolve_requester(self._req()))

    def test_player_and_unknown_claims_fail_closed_to_none(self):
        self.assertIsNone(serve._resolve_requester(self._req('requesterId=player')))
        self.assertIsNone(serve._resolve_requester(self._req('requesterId=unknown')))

    def test_bad_key_returns_none(self):
        with unittest.mock.patch.object(serve, 'verify_agent_key', return_value=False):
            self.assertIsNone(serve._resolve_requester(self._req('requesterId=ada')))

    def test_valid_key_returns_claimed_id(self):
        with unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            self.assertEqual(serve._resolve_requester(self._req('requesterId=ada')), 'ada')


class AgentFilesEndpoints(unittest.TestCase):
    def setUp(self):
        self.agent = _uid('ag')
        self.base = os.path.join(serve.AGENTS_DIR, self.agent)
        os.makedirs(os.path.join(self.base, 'notes'), exist_ok=True)
        os.makedirs(os.path.join(self.base, 'reports'), exist_ok=True)
        os.makedirs(os.path.join(self.base, 'conversations'), exist_ok=True)
        os.makedirs(os.path.join(self.base, '.git'), exist_ok=True)
        os.makedirs(os.path.join(self.base, 'sub'), exist_ok=True)
        for rel in ('notes/a.md', 'reports/r.md', 'sub/b.md', '.hidden.md', 'conversations/c.json', '.git/config', '.DS_Store'):
            with open(os.path.join(self.base, rel), 'w') as f:
                f.write(rel)

    def test_unknown_agent_returns_empty_list(self):
        r = _client().get('/api/agent-files', params={'agentId': 'no-such-agent'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json(), {'files': []})

    def test_lists_visible_files_and_hides_skips(self):
        r = _client().get('/api/agent-files', params={'agentId': self.agent})
        self.assertEqual(r.status_code, 200)
        paths = [f['path'] for f in r.json()['files']]
        self.assertIn('notes/a.md', paths)
        self.assertIn('reports/r.md', paths)
        self.assertIn('sub/b.md', paths)
        self.assertNotIn('conversations/c.json', paths)
        self.assertNotIn('.hidden.md', paths)
        self.assertNotIn('.git/config', paths)
        self.assertNotIn('.DS_Store', paths)
        self.assertEqual(paths, sorted(paths))

    def test_own_reports_are_hidden_from_the_subject(self):
        with unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().get('/api/agent-files', params={'agentId': self.agent, 'requesterId': self.agent})
        self.assertEqual(r.status_code, 200)
        paths = [f['path'] for f in r.json()['files']]
        self.assertIn('notes/a.md', paths)
        self.assertNotIn('reports/r.md', paths)

    def test_read_traversal_is_rejected(self):
        r = _client().get('/api/agent-files/read', params={'agentId': self.agent, 'path': '../../etc/passwd'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('invalid path', r.json()['error'])

    def test_read_invisible_path_is_403(self):
        r = _client().get('/api/agent-files/read', params={'agentId': self.agent, 'path': '.hidden.md'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('not readable', r.json()['error'])

    def test_read_own_report_is_403(self):
        with unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().get('/api/agent-files/read', params={'agentId': self.agent, 'path': 'reports/r.md', 'requesterId': self.agent})
        self.assertEqual(r.status_code, 403)

    def test_read_missing_file_is_404(self):
        r = _client().get('/api/agent-files/read', params={'agentId': self.agent, 'path': 'notes/zz.md'})
        self.assertEqual(r.status_code, 404)

    def test_read_oversized_file_is_413(self):
        with open(os.path.join(self.base, 'notes/huge.md'), 'w') as f:
            f.write('x' * 250_000)
        r = _client().get('/api/agent-files/read', params={'agentId': self.agent, 'path': 'notes/huge.md'})
        self.assertEqual(r.status_code, 413)

    def test_read_success(self):
        r = _client().get('/api/agent-files/read', params={'agentId': self.agent, 'path': 'notes/a.md'})
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body['agentId'], self.agent)
        self.assertEqual(body['path'], 'notes/a.md')
        self.assertEqual(body['content'], 'notes/a.md')


class BrowseEndpoint(unittest.TestCase):
    def _post(self, body, **patches):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'):
            return _client().post('/api/browse', json=body)

    def test_disabled_browsing_returns_403(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'BROWSING_ENABLED', False):
            r = _client().post('/api/browse', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 403)

    def test_missing_key_returns_500(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None):
            r = _client().post('/api/browse', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 500)

    def test_budget_exhausted_blocks(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, '_page_budget_exhausted', return_value=True):
            r = _client().post('/api/browse', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/browse', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_missing_url_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().post('/api/browse', json={'agentId': 'ada'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('url is required', r.json()['error'])

    def test_invalid_scheme_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/browse', json={'url': 'ftp://example.com/x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'invalid or non-http(s) scheme')

    def test_private_host_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/browse', json={'url': 'http://localhost/x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'private/internal host')

    def test_jev_denial_blocks(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=False), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False):
            r = _client().post('/api/browse', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'research'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])

    def test_allowlist_path_success_non_render(self):
        body = {'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'check the weather'}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=True), \
             unittest.mock.patch.object(serve, '_fetch_page_sync',
                                        return_value=('https://example.com/x', 'text/html', '<html><body>weather report</body></html>', False, None)), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/browse', json=body)
        self.assertEqual(r.status_code, 200, r.text)
        resp = r.json()
        self.assertTrue(resp['allowed'])
        self.assertEqual(resp['text'], 'weather report')
        self.assertIn('EXTERNAL_DATA', resp['textForModel'])
        self.assertEqual(log.call_args_list[0][0][2]['decision'], 'allowed_by_allowlist')

    def test_jev_allow_success_with_render(self):
        long_text = ('RENDERED ' * 5000)  # > 20000 chars -> truncated True
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=False), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev(trace='trace-browse')), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, 'sync_playwright', object()), \
             unittest.mock.patch.object(serve, '_fetch_rendered_page_sync',
                                        return_value=('https://example.com/rendered', long_text, [])), \
             unittest.mock.patch.object(serve, 'record_browse_success') as trail:
            r = _client().post('/api/browse', json={'url': 'http://example.com/app', 'agentId': 'ada', 'purpose': 'research', 'render': True})
        self.assertEqual(r.status_code, 200, r.text)
        resp = r.json()
        self.assertTrue(resp['allowed'])
        self.assertTrue(resp['truncated'])
        self.assertEqual(resp['url'], 'https://example.com/rendered')
        trail.assert_called_once()

    def test_render_without_playwright_reports_fetch_error(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=True), \
             unittest.mock.patch.object(serve, 'sync_playwright', None):
            r = _client().post('/api/browse', json={'url': 'http://example.com/app', 'agentId': 'ada', 'purpose': 'x', 'render': True})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['allowed'])
        self.assertIn('playwright is not installed', r.json()['error'])

    def test_visual_success_captures_screenshot(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=True), \
             unittest.mock.patch.object(serve, '_fetch_page_sync',
                                        return_value=('https://example.com/x', 'text/html', '<html>visual</html>', False, None)), \
             unittest.mock.patch.object(serve, '_find_chrome', return_value='/usr/bin/google-chrome'), \
             unittest.mock.patch.object(serve, '_screenshot_url_sync', return_value='BASE64PNG'):
            r = _client().post('/api/browse', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'ui', 'visual': True})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['imageBase64'], 'BASE64PNG')

    def test_visual_failure_is_best_effort(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=True), \
             unittest.mock.patch.object(serve, '_fetch_page_sync',
                                        return_value=('https://example.com/x', 'text/html', '<html>visual</html>', False, None)), \
             unittest.mock.patch.object(serve, '_find_chrome', return_value='/usr/bin/google-chrome'), \
             unittest.mock.patch.object(serve, '_screenshot_url_sync', side_effect=Exception('chrome died')):
            r = _client().post('/api/browse', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'ui', 'visual': True})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIsNone(r.json()['imageBase64'])

    def test_fetch_failure_is_reported(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=True), \
             unittest.mock.patch.object(serve, '_fetch_page_sync', side_effect=ConnectionError('network down')):
            r = _client().post('/api/browse', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'x'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['allowed'])
        self.assertIn('network down', r.json()['error'])


class AllowlistRequestEndpoint(unittest.TestCase):
    def _post(self, body, **patches):
        return _client().post('/api/allowlist/request', json=body)

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada', 'host': 'example.com'})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_missing_target_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('host or url is required', r.json()['error'])

    def test_unreadable_hostname_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada', 'url': ':'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('hostname', r.json()['error'])

    def test_private_host_is_refused(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=False):
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada', 'host': 'internal.example.com'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('private or internal', r.json()['error'])

    def test_already_allowlisted_host(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=True):
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada', 'host': 'example.com'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['allowlisted'])

    def test_duplicate_pending_request_is_reported(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=False), \
             unittest.mock.patch.object(serve, '_load_escalations',
                                        return_value={'esc-1': {'kind': 'allowlist request', 'status': 'pending', 'note': 'example.com'}}):
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada', 'host': 'https://example.com/page'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['requested'])
        self.assertEqual(r.json()['escalationId'], 'esc-1')

    def test_success_files_escalation(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_is_allowlisted_host', return_value=False), \
             unittest.mock.patch.object(serve, '_load_escalations', return_value={}), \
             unittest.mock.patch.object(serve, 'create_escalation', return_value='esc-9'), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/allowlist/request', json={'agentId': 'ada', 'host': 'api.example.com', 'purpose': 'crawl the docs'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()['requested'])
        self.assertEqual(r.json()['escalationId'], 'esc-9')
        self.assertEqual(log.call_args[0][1], 'allowlist_request')


class RoomGates(unittest.TestCase):
    def test_weatherstation_live_room_override_allows(self):
        self.assertTrue(serve._agent_is_in_weatherstation('eli', 'weatherstation'))

    def test_weatherstation_live_room_mismatch_checks_state(self):
        with unittest.mock.patch.object(serve, 'get_state_from_db', return_value={'agents': {'eli': {'inRoom': 'bank'}}}):
            self.assertFalse(serve._agent_is_in_weatherstation('eli', 'lobby'))

    def test_sandbox_room_live_room_override_allows(self):
        self.assertTrue(serve._agent_is_in_sandbox_room('dev', 'observatory'))
        self.assertTrue(serve._agent_is_in_sandbox_room('dev', 'pressoffice'))

    def test_sandbox_room_live_room_mismatch_checks_state(self):
        with unittest.mock.patch.object(serve, 'get_state_from_db', return_value={'agents': {'dev': {'inRoom': 'bank'}}}):
            self.assertFalse(serve._agent_is_in_sandbox_room('dev', 'hangout'))

    def test_sandbox_room_state_in_room_allows(self):
        with unittest.mock.patch.object(serve, 'get_state_from_db', return_value={'agents': {'dev': {'inRoom': 'observatory'}}}):
            self.assertTrue(serve._agent_is_in_sandbox_room('dev'))


class CurlRequestSync(unittest.TestCase):
    def _run(self, resp, headers=None, body=None, host_ok=True):
        with unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=host_ok), \
             unittest.mock.patch.object(serve.urllib.request, 'urlopen', return_value=resp):
            return serve._curl_request_sync('GET', 'https://api.example.com/x', headers or {}, body)

    def test_success_without_user_agent_adds_one(self):
        resp = _FakeResp('https://api.example.com/x', {'Content-Type': 'application/json'}, 200, b'{"ok":true}')
        result = self._run(resp)
        self.assertEqual(result['status'], 200)
        self.assertEqual(result['finalUrl'], 'https://api.example.com/x')
        self.assertEqual(result['headers']['Content-Type'], 'application/json')
        self.assertEqual(result['body'], '{"ok":true}')
        self.assertFalse(result['truncated'])

    def test_body_is_encoded(self):
        resp = _FakeResp('https://api.example.com/x', {}, 200, b'accepted')
        with unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve.urllib.request, 'urlopen', return_value=resp) as urlopen:
            serve._curl_request_sync('POST', 'https://api.example.com/x', {'Content-Type': 'text/plain'}, 'payload')
        req = urlopen.call_args[0][0]
        self.assertEqual(req.data, b'payload')
        self.assertEqual(req.method, 'POST')
        self.assertEqual(req.headers.get('User-agent'), 'AIThinkTankAgent/1.0')

    def test_existing_user_agent_is_kept(self):
        resp = _FakeResp('https://api.example.com/x', {}, 200, b'ok')
        with unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve.urllib.request, 'urlopen', return_value=resp) as urlopen:
            serve._curl_request_sync('GET', 'https://api.example.com/x', {'User-Agent': 'custom/1.0'}, None)
        req = urlopen.call_args[0][0]
        self.assertEqual(req.headers.get('User-agent'), 'custom/1.0')

    def test_redirect_to_disallowed_host_raises(self):
        resp = _FakeResp('http://10.0.0.9/x', {}, 200, b'x')
        with self.assertRaises(ValueError):
            self._run(resp, host_ok=False)

    def test_truncation_flag(self):
        resp = _FakeResp('https://api.example.com/x', {}, 200, b'x' * (serve.CURL_MAX_BODY_BYTES + 100))
        result = self._run(resp)
        self.assertTrue(result['truncated'])
        self.assertEqual(len(result['body']), serve.CURL_MAX_BODY_BYTES)


class CurlEndpoint(unittest.TestCase):
    def _post(self, body, **patches):
        return _client().post('/api/curl', json=body)

    def _ok_patches(self, agent):
        return [
            unittest.mock.patch.object(serve, 'verify_session', return_value=True),
            unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'),
            unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True),
            unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True),
            unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False),
            unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True),
            unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()),
            unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True),
        ]

    def test_disabled_browsing_returns_403(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'BROWSING_ENABLED', False):
            r = _client().post('/api/curl', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 403)

    def test_missing_key_returns_500(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None):
            r = _client().post('/api/curl', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 500)

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_not_in_weatherstation_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=False), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'not in the Weather Station and no active temporary grant')

    def test_missing_url_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False):
            r = _client().post('/api/curl', json={'agentId': 'ada'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('url is required', r.json()['error'])

    def test_invalid_method_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False):
            r = _client().post('/api/curl', json={'url': 'http://x', 'agentId': 'ada', 'method': 'TRACE'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('method must be one of', r.json()['error'])

    def test_invalid_scheme_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'ftp://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'invalid or non-http(s) scheme')

    def test_private_host_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'http://localhost/x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'private/internal host')

    def test_jev_denial_blocks(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False):
            r = _client().post('/api/curl', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'api'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])

    def test_invalid_capability_handle_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, 'resolve_capability_handle', return_value=None), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'http://example.com/x', 'agentId': 'ada', 'capabilityHandle': 'bad'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'capability handle invalid, expired, or out of scope')

    def test_fetch_failure_is_reported(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_curl_request_sync', side_effect=ConnectionError('refused')), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'api'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['allowed'])
        self.assertIn('refused', r.json()['error'])
        self.assertEqual(log.call_args[0][2]['decision'], 'allowed_but_failed')

    def test_success_without_handle(self):
        result = {'status': 200, 'finalUrl': 'https://example.com/x', 'headers': {'Content-Type': 'text/plain'}, 'body': 'hello', 'truncated': False}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_curl_request_sync', return_value=result):
            r = _client().post('/api/curl', json={'url': 'http://example.com/x', 'agentId': 'ada', 'purpose': 'api'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()['allowed'])
        self.assertEqual(r.json()['body'], 'hello')

    def test_success_with_handle_redacts_secret(self):
        result = {'status': 200, 'finalUrl': 'https://api.example.com/x',
                  'headers': {'Content-Type': 'application/json', 'X-Echo': 'TOP-SECRET'},
                  'body': '{"token": "TOP-SECRET"}', 'truncated': False}
        grant = {'credential_name': 'api', 'purpose': 'scoped use', 'service': 'svc', 'secret': 'TOP-SECRET'}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_weatherstation', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, 'resolve_capability_handle', return_value=grant), \
             unittest.mock.patch.object(serve, '_curl_request_sync', return_value=result), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/curl', json={'url': 'http://api.example.com/x', 'agentId': 'ada', 'capabilityHandle': 'h1'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertNotIn('TOP-SECRET', r.text, 'the injected credential must never reach the caller')
        self.assertIn('[REDACTED]', r.json()['body'])
        self.assertEqual(log.call_args[0][2]['decision'], 'allowed')


class SandboxDownloadEndpoint(unittest.TestCase):
    def _post(self, body, **patches):
        return _client().post('/api/sandbox-download', json=body)

    def _ok_patches(self):
        return [
            unittest.mock.patch.object(serve, 'verify_session', return_value=True),
            unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'),
            unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True),
            unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True),
            unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False),
            unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True),
            unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()),
            unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True),
        ]

    def test_disabled_returns_403(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'BROWSING_ENABLED', False):
            r = _client().post('/api/sandbox-download', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 403)

    def test_missing_key_returns_500(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None):
            r = _client().post('/api/sandbox-download', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 500)

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-download', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_not_in_sandbox_room_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=False), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-download', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['decision'], 'blocked')

    def test_missing_fields_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False):
            r = _client().post('/api/sandbox-download', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('url, filename, and sandboxId are required', r.json()['error'])

    def test_bad_filename_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False):
            r = _client().post('/api/sandbox-download', json={'url': 'http://x', 'agentId': 'ada', 'filename': '///', 'sandboxId': 's1'})
        self.assertEqual(r.status_code, 400)

    def test_invalid_scheme_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-download', json={'url': 'ftp://x', 'agentId': 'ada', 'filename': 'x.bin', 'sandboxId': 's1'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'invalid or non-http(s) scheme')

    def test_private_host_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-download', json={'url': 'http://localhost/x', 'agentId': 'ada', 'filename': 'x.bin', 'sandboxId': 's1'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'private/internal host')

    def test_jev_denial_blocks(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False):
            r = _client().post('/api/sandbox-download', json={'url': 'http://example.com/x', 'agentId': 'ada', 'filename': 'x.bin', 'sandboxId': 's1'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])

    def test_fetch_failure_is_reported(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync', side_effect=ValueError('dl failed')):
            r = _client().post('/api/sandbox-download', json={'url': 'http://example.com/x', 'agentId': 'ada', 'filename': 'x.bin', 'sandboxId': 's1'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['allowed'])
        self.assertFalse(r.json()['ok'])
        self.assertIn('dl failed', r.json()['reason'])

    def test_success_writes_into_sandbox(self):
        sandbox_id = _uid('sbox')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev(trace='trace-sd')), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, '_download_file_sync',
                                        return_value=('https://example.com/data.csv', 'text/csv', b'csvdata', False)), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-download', json={'url': 'http://example.com/data.csv', 'agentId': 'ada', 'filename': 'data.csv', 'sandboxId': sandbox_id})
        self.assertEqual(r.status_code, 200, r.text)
        resp = r.json()
        self.assertTrue(resp['allowed'])
        self.assertTrue(resp['ok'])
        self.assertEqual(resp['path'], 'downloads/data.csv')
        sandbox_dir = serve._sandbox_dir_for(sandbox_id)
        with open(os.path.join(sandbox_dir, 'downloads', 'data.csv'), 'rb') as f:
            self.assertEqual(f.read(), b'csvdata')
        self.assertEqual(log.call_args[0][2]['decision'], 'allowed')


class SandboxSavePageEndpoint(unittest.TestCase):
    def test_disabled_returns_403(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'BROWSING_ENABLED', False):
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 403)

    def test_missing_key_returns_500(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None):
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://x'})
        self.assertEqual(r.status_code, 500)

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_not_in_room_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=False), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://x', 'agentId': 'ada'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['decision'], 'blocked')

    def test_missing_fields_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False):
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://x', 'agentId': 'ada', 'sandboxId': 's1'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('url, filename, sandboxId, and content are required', r.json()['error'])

    def test_bad_filename_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False):
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://x', 'agentId': 'ada', 'filename': '///', 'sandboxId': 's1', 'content': 'c'})
        self.assertEqual(r.status_code, 400)

    def test_invalid_scheme_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-save-page', json={'url': 'ftp://x', 'agentId': 'ada', 'filename': 'p.html', 'sandboxId': 's1', 'content': 'c'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'invalid or non-http(s) scheme')

    def test_private_host_is_blocked(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://localhost/x', 'agentId': 'ada', 'filename': 'p.html', 'sandboxId': 's1', 'content': 'c'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])
        self.assertEqual(log.call_args[0][2]['reason'], 'private/internal host')

    def test_jev_denial_blocks(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False):
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://example.com/x', 'agentId': 'ada', 'filename': 'p.html', 'sandboxId': 's1', 'content': 'c'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['allowed'])

    def test_success_writes_and_truncates(self):
        sandbox_id = _uid('sbox')
        content = 'x' * (serve.SANDBOX_SAVE_MAX_BYTES + 50)
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev(trace='trace-ssp')), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://example.com/page', 'agentId': 'ada', 'filename': 'page.html', 'sandboxId': sandbox_id, 'content': content})
        self.assertEqual(r.status_code, 200, r.text)
        resp = r.json()
        self.assertTrue(resp['allowed'])
        self.assertTrue(resp['ok'])
        self.assertTrue(resp['truncated'])
        sandbox_dir = serve._sandbox_dir_for(sandbox_id)
        with open(os.path.join(sandbox_dir, 'downloads', 'page.html'), 'rb') as f:
            self.assertEqual(len(f.read()), serve.SANDBOX_SAVE_MAX_BYTES)
        self.assertEqual(log.call_args[0][2]['decision'], 'allowed')

    def test_success_short_content_not_truncated(self):
        sandbox_id = _uid('sbox')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True), \
             unittest.mock.patch.object(serve, '_agent_is_in_sandbox_room', return_value=True), \
             unittest.mock.patch.object(serve, '_has_active_temp_access', return_value=False), \
             unittest.mock.patch.object(serve, '_is_safe_public_host', return_value=True), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True):
            r = _client().post('/api/sandbox-save-page', json={'url': 'http://example.com/page', 'agentId': 'ada', 'filename': 'page.html', 'sandboxId': sandbox_id, 'content': 'short page text'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertFalse(r.json()['truncated'])


class AccessRequestEndpoint(unittest.TestCase):
    def _post(self, body, **patches):
        return _client().post('/api/access/request', json=body)

    def test_rate_limited_returns_429(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'check_rate_limit', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/access/request', json={'agentId': 'ada', 'capability': 'curl', 'reason': 'x'})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(log.call_args[0][1], 'rate_limited')

    def test_invalid_capability_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/access/request', json={'agentId': 'ada', 'capability': 'nope', 'reason': 'x'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('capability must be one of', r.json()['error'])

    def test_missing_reason_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/access/request', json={'agentId': 'ada', 'capability': 'curl'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('reason is required', r.json()['error'])

    def test_task_id_mismatch_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'get_state_from_db',
                                        return_value={'agents': {'ada': {'task': 'other-task'}}}):
            r = _client().post('/api/access/request', json={'agentId': 'ada', 'capability': 'curl', 'reason': 'x', 'taskId': 't1'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('currently working on', r.json()['error'])

    def test_missing_key_returns_500(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None):
            r = _client().post('/api/access/request', json={'agentId': 'ada', 'capability': 'curl', 'reason': 'x'})
        self.assertEqual(r.status_code, 500)

    def test_jev_denial_denies(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/access/request', json={'agentId': 'ada', 'capability': 'curl', 'reason': 'legit need', 'supervisorId': 'faye'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()['approved'])
        self.assertEqual(log.call_args[0][1], 'access_request')

    def test_success_grants_temporary_access(self):
        agent = _uid('acc')
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev(trace='trace-ar')), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True), \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/access/request', json={'agentId': agent, 'capability': 'curl', 'reason': 'need to check a weather API', 'supervisorId': 'faye'})
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertTrue(body['approved'])
        self.assertFalse(body['untilStoryComplete'])
        self.assertGreater(body['expiresAt'], time.time())
        self.assertTrue(serve._has_active_temp_access(agent, 'curl'))
        self.assertEqual(log.call_args[0][2]['decision'], 'approved')

    def test_success_story_scoped_grant(self):
        agent = _uid('acc')
        state = {'agents': {agent: {'task': 't1'}}}
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, 'get_state_from_db', return_value=state), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True):
            r = _client().post('/api/access/request', json={'agentId': agent, 'capability': 'curl', 'reason': 'for the story', 'supervisorId': 'faye', 'taskId': 't1'})
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertTrue(body['approved'])
        self.assertTrue(body['untilStoryComplete'])
        self.assertIsNone(body['expiresAt'])
        self.assertTrue(serve._has_active_temp_access(agent, 'curl'))


class ClassifyCommand(unittest.TestCase):
    def test_missing_key_fails_closed(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', None):
            allowed, reason = asyncio.run(serve._classify_command('echo hi', 'test', 'ada'))
        self.assertFalse(allowed)
        self.assertEqual(reason, 'OPENROUTER_API_KEY not set')

    def test_confident_allow_passes(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev()), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=True):
            allowed, reason = asyncio.run(serve._classify_command('echo hi', 'runs a test', 'ada'))
        self.assertTrue(allowed)
        self.assertEqual(reason, 'allow')

    def test_block_blocks(self):
        with unittest.mock.patch.object(serve, 'OPENROUTER_API_KEY', 'k'), \
             unittest.mock.patch.object(serve, '_jev_quorum_decision', _jev(decision='block', confidence=0.99)), \
             unittest.mock.patch.object(serve, '_jev_safety_gate', return_value=False):
            allowed, reason = asyncio.run(serve._classify_command('rm -rf /', 'cleanup', 'ada'))
        self.assertFalse(allowed)
        self.assertEqual(reason, 'block')


class CredentialVaultEndpoints(unittest.TestCase):
    def setUp(self):
        if os.path.exists(serve.PASSPORT_PATH):
            os.remove(serve.PASSPORT_PATH)

    def test_list_credentials_blocks_agents(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().get('/api/keys/credentials', params={'requesterId': 'ada'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('player-only', r.json()['error'])

    def test_list_credentials_returns_entries(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_list_credentials',
                                        return_value=[{'name': 'api', 'service': 'svc'}]):
            r = _client().get('/api/keys/credentials')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['credentials'], [{'name': 'api', 'service': 'svc'}])

    def test_add_credential_blocks_agents(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().post('/api/keys/credentials', json={'name': 'x', 'value': 'v'}, params={'requesterId': 'ada'})
        self.assertEqual(r.status_code, 403)

    def test_add_credential_missing_fields_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/keys/credentials', json={'name': 'x'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('name and value are required', r.json()['error'])

    def test_add_credential_bad_name_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/keys/credentials', json={'name': 'bad name!', 'value': 'v'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('must be [a-zA-Z0-9_-]', r.json()['error'])

    def test_add_credential_success_stores_and_chains(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_store_credential') as store, \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/keys/credentials', json={'name': 'api', 'service': 'my-svc', 'value': 'secret-value'})
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertTrue(body['ok'])
        self.assertEqual(body['credential'], {'name': 'api', 'service': 'my-svc'})
        store.assert_called_once_with('api', 'my-svc', 'secret-value')
        self.assertEqual(log.call_args[0][1], 'credential_stored')
        data = serve._load_passport()
        self.assertEqual(data['blocks'][0]['kind'], 'credential_stored')

    def test_delete_credential_blocks_agents(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'verify_agent_key', return_value=True):
            r = _client().delete('/api/keys/credentials/api', params={'requesterId': 'ada'})
        self.assertEqual(r.status_code, 403)

    def test_delete_credential_success(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, '_delete_credential') as delete, \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().delete('/api/keys/credentials/api')
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()['ok'])
        delete.assert_called_once_with('api')
        self.assertEqual(log.call_args[0][1], 'credential_deleted')
        data = serve._load_passport()
        self.assertEqual(data['blocks'][0]['kind'], 'credential_deleted')


class CapabilityHandles(unittest.TestCase):
    def setUp(self):
        if os.path.exists(serve.PASSPORT_PATH):
            os.remove(serve.PASSPORT_PATH)

    def _post_handle(self, body, **patches):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            return _client().post('/api/keys/handles', json=body)

    def test_requires_session(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=False):
            r = _client().post('/api/keys/handles', json={})
        self.assertEqual(r.status_code, 401)

    def test_missing_fields_returns_400(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True):
            r = _client().post('/api/keys/handles', json={'agentId': 'ada'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('agentId, credentialName, and purpose are required', r.json()['error'])

    def test_string_allowed_methods_is_parsed(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'mint_capability_handle',
                                        return_value=('handle-1', None)) as mint:
            r = _client().post('/api/keys/handles', json={'agentId': 'ada', 'credentialName': 'api', 'purpose': 'p', 'allowedMethods': 'GET, POST'})
        self.assertEqual(r.status_code, 200, r.text)
        args = mint.call_args[0]
        self.assertEqual(args[4], ['GET', 'POST'])

    def test_string_allowed_hosts_is_parsed(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'mint_capability_handle',
                                        return_value=('handle-1', None)) as mint:
            r = _client().post('/api/keys/handles', json={'agentId': 'ada', 'credentialName': 'api', 'purpose': 'p', 'allowedHosts': 'api.a.com,api.b.com'})
        self.assertEqual(r.status_code, 200, r.text)
        args = mint.call_args[0]
        self.assertEqual(args[3], ['api.a.com', 'api.b.com'])

    def test_wildcard_host_stays_star(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'mint_capability_handle',
                                        return_value=('handle-1', None)) as mint:
            r = _client().post('/api/keys/handles', json={'agentId': 'ada', 'credentialName': 'api', 'purpose': 'p', 'allowedHosts': '*'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(mint.call_args[0][3], ['*'])

    def test_unknown_credential_refusal_is_404(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'mint_capability_handle', return_value=(None, 'unknown credential')):
            r = _client().post('/api/keys/handles', json={'agentId': 'ada', 'credentialName': 'ghost', 'purpose': 'p'})
        self.assertEqual(r.status_code, 404)
        self.assertIn('unknown credential', r.json()['error'])

    def test_other_refusal_is_403(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'mint_capability_handle', return_value=(None, 'over cap')):
            r = _client().post('/api/keys/handles', json={'agentId': 'ada', 'credentialName': 'api', 'purpose': 'p'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(r.json()['error'], 'over cap')

    def test_success_mints_handle(self):
        with unittest.mock.patch.object(serve, 'verify_session', return_value=True), \
             unittest.mock.patch.object(serve, 'mint_capability_handle',
                                        return_value=('minted-nonce', None)) as mint, \
             unittest.mock.patch.object(serve, 'log_action') as log:
            r = _client().post('/api/keys/handles', json={'agentId': 'ada', 'credentialName': 'api', 'purpose': 'scoped use', 'ttlSec': 600})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['handle'], 'minted-nonce')
        self.assertEqual(mint.call_args[0][1], 'api')
        self.assertEqual(log.call_args[0][1], 'handle_minted')
        data = serve._load_passport()
        self.assertEqual(data['blocks'][0]['kind'], 'handle_minted')

    def _insert_cred(self, name, service='svc'):
        with serve._db() as conn:
            conn.execute('INSERT OR REPLACE INTO external_credentials (name, service, encrypted_value, created_at) VALUES (?, ?, ?, ?)',
                         (name, service, 'encrypted-token', time.time()))

    def _insert_handle(self, handle, agent_id, cred='api', hosts='["api.example.com"]', methods='["GET"]', expires_at=None):
        with serve._db() as conn:
            conn.execute('INSERT INTO capability_handles (handle, agent_id, credential_name, purpose, allowed_hosts, allowed_methods, granted_by, expires_at, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
                         (handle, agent_id, cred, 'purpose', hosts, methods, 'player', time.time() + 3600 if expires_at is None else expires_at, time.time()))

    def test_resolve_empty_inputs_is_none(self):
        self.assertIsNone(serve.resolve_capability_handle('', 'h', 'GET', 'http://x'))

    def test_resolve_no_row_is_none(self):
        self.assertIsNone(serve.resolve_capability_handle('ada', 'no-such-handle', 'GET', 'http://api.example.com'))

    def test_resolve_digitalocean_disabled_is_none(self):
        self._insert_cred('digitalocean')
        self._insert_handle('do-handle', 'ada', cred='digitalocean')
        with unittest.mock.patch.object(serve, '_digitalocean_enabled', return_value=False):
            self.assertIsNone(serve.resolve_capability_handle('ada', 'do-handle', 'GET', 'http://api.example.com'))

    def test_resolve_bad_json_scope_is_none(self):
        self._insert_cred('api')
        self._insert_handle('bad-json-handle', 'ada', hosts='not-json')
        self.assertIsNone(serve.resolve_capability_handle('ada', 'bad-json-handle', 'GET', 'http://api.example.com'))

    def test_resolve_expired_handle_is_none_and_deleted(self):
        self._insert_cred('api')
        self._insert_handle('expired-handle', 'ada', expires_at=time.time() - 100)
        self.assertIsNone(serve.resolve_capability_handle('ada', 'expired-handle', 'GET', 'http://api.example.com'))
        with serve._db() as conn:
            row = conn.execute('SELECT 1 FROM capability_handles WHERE handle = ?', ('expired-handle',)).fetchone()
        self.assertIsNone(row, 'an expired handle is deleted on resolve')

    def test_resolve_url_without_host_is_none(self):
        self._insert_cred('api')
        self._insert_handle('no-host-handle', 'ada')
        self.assertIsNone(serve.resolve_capability_handle('ada', 'no-host-handle', 'GET', 'not-a-url'))

    def test_resolve_host_out_of_scope_is_none(self):
        self._insert_cred('api')
        self._insert_handle('host-handle', 'ada')
        self.assertIsNone(serve.resolve_capability_handle('ada', 'host-handle', 'GET', 'http://evil.example.net'))

    def test_resolve_method_out_of_scope_is_none(self):
        self._insert_cred('api')
        self._insert_handle('method-handle', 'ada', methods='["GET"]')
        self.assertIsNone(serve.resolve_capability_handle('ada', 'method-handle', 'POST', 'http://api.example.com'))

    def test_resolve_missing_credential_row_is_none(self):
        self._insert_handle('no-cred-handle', 'ada', cred='ghost-cred')
        self.assertIsNone(serve.resolve_capability_handle('ada', 'no-cred-handle', 'GET', 'http://api.example.com'))

    def test_resolve_secret_unreadable_is_none(self):
        self._insert_cred('api')
        self._insert_handle('secret-handle', 'ada')
        with unittest.mock.patch.object(serve, '_open_secret', return_value=None):
            self.assertIsNone(serve.resolve_capability_handle('ada', 'secret-handle', 'GET', 'http://api.example.com'))

    def test_resolve_success_returns_scoped_headers(self):
        self._insert_cred('api', service='github')
        self._insert_handle('ok-handle', 'ada')
        with unittest.mock.patch.object(serve, '_open_secret', return_value='decrypted-token'):
            result = serve.resolve_capability_handle('ada', 'ok-handle', 'GET', 'http://api.example.com/repos')
        self.assertEqual(result['credential_name'], 'api')
        self.assertEqual(result['service'], 'github')
        self.assertEqual(result['secret'], 'decrypted-token')
        self.assertEqual(result['purpose'], 'purpose')

    def test_resolve_wildcard_host_success(self):
        self._insert_cred('api')
        self._insert_handle('wild-handle', 'ada', hosts='"*"')
        with unittest.mock.patch.object(serve, '_open_secret', return_value='t'):
            result = serve.resolve_capability_handle('ada', 'wild-handle', 'GET', 'http://anything.example.com')
        self.assertIsNotNone(result)


if __name__ == '__main__':
    unittest.main()